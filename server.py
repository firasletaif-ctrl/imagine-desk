#!/usr/bin/env python3
"""Imagine Events Desk — prototype multi-événements, multi-postes.

Utilise la bibliothèque standard Python et openpyxl pour les classeurs .xlsx. Pour une mise en production,
prévoir HTTPS, sauvegardes et une base PostgreSQL administrée.
"""
from __future__ import annotations

import json
import os
import re
import secrets
import sqlite3
import uuid
import unicodedata
import hashlib
import hmac
import time
import tempfile
from datetime import datetime, timezone, timedelta
from email import policy
from email.parser import BytesParser
from http.cookies import SimpleCookie
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

try:  # PostgreSQL n’est nécessaire que si DATABASE_URL est fourni (Render, OVH…).
    import psycopg2
    from psycopg2.extras import RealDictCursor
except Exception:  # pragma: no cover - l’application reste utilisable en SQLite
    psycopg2 = None
    RealDictCursor = None

DB_INTEGRITY_ERRORS = (sqlite3.IntegrityError,)
if psycopg2 is not None:
    DB_INTEGRITY_ERRORS += (psycopg2.IntegrityError,)

ROOT = Path(__file__).resolve().parent
PUBLIC = ROOT / "public"
DB_PATH = Path(os.environ.get("IMAGINE_DB_PATH", ROOT / "data" / "imagine_desk.sqlite3"))
# PostgreSQL managé (Render, OVH, Neon…) : une seule variable DATABASE_URL suffit.
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = "postgresql://" + DATABASE_URL[len("postgres://"):]
USE_POSTGRES = DATABASE_URL.startswith("postgresql://")
HOST = os.environ.get("HOST", "0.0.0.0")
PORT = int(os.environ.get("PORT", "8000"))
MAX_BODY = 12 * 1024 * 1024
MAX_BACKUP_FILE = 50 * 1024 * 1024
MAX_BACKUP_REQUEST = MAX_BACKUP_FILE + 1024 * 1024
CATEGORIES = {"participant", "vip", "speaker", "press", "team"}
PAYMENT_STATUSES = {"unpaid", "paid", "exempt"}
PAYMENT_METHODS = {"", "cash", "card", "transfer", "cheque", "mobile", "other"}
MAX_PAYMENT_AMOUNT = 10_000_000
DEFAULT_CURRENCY = "TND"
SESSION_COOKIE = "imagine_events_session"
SESSION_TTL_SECONDS = 12 * 60 * 60
PASSWORD_ITERATIONS = 310_000
SETUP_TOKEN = os.environ.get("IMAGINE_SETUP_TOKEN", "").strip()
EPHEMERAL_PREFIXES = ("/tmp/", "/var/tmp/", "/dev/shm/", "/run/")
LOGIN_FAILURES = {}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def db_display_path() -> str:
    """Chemin lisible de la base (relatif au projet quand c’est possible)."""
    if USE_POSTGRES:
        return "PostgreSQL (DATABASE_URL)"
    try:
        return str(DB_PATH.resolve().relative_to(ROOT.resolve()))
    except (ValueError, OSError):
        return str(DB_PATH)


def db_is_ephemeral() -> bool:
    """Vrai si la base est sur un dossier temporaire, perdu au redémarrage du serveur."""
    if USE_POSTGRES:
        return False
    try:
        resolved = str(DB_PATH.resolve())
    except OSError:
        resolved = str(DB_PATH)
    return resolved.startswith(EPHEMERAL_PREFIXES)


def db_backend() -> str:
    return "postgres" if USE_POSTGRES else "sqlite"


def redact_database_url(url: str = "") -> str:
    """Masque le mot de passe avant tout affichage ou journal."""
    url = url or DATABASE_URL
    return re.sub(r"://([^:/@]+):[^@]*@", r"://\1:***@", url)


def storage_info() -> dict:
    if USE_POSTGRES:
        host = ""
        try:
            from urllib.parse import urlsplit
            host = urlsplit(DATABASE_URL).hostname or ""
        except ValueError:
            host = ""
        return {
            "backend": "postgres",
            "path": f"PostgreSQL · {host}" if host else "PostgreSQL (DATABASE_URL)",
            "exists": True,
            "sizeBytes": 0,
            "ephemeral": False,
            "envOverride": True,
            "warnings": "",
            "detail": redact_database_url(),
        }
    exists = DB_PATH.is_file()
    try:
        size = DB_PATH.stat().st_size if exists else 0
    except OSError:
        size = 0
    return {
        "backend": "sqlite",
        "path": db_display_path(),
        "exists": exists,
        "sizeBytes": size,
        "ephemeral": db_is_ephemeral(),
        "envOverride": bool(os.environ.get("IMAGINE_DB_PATH", "").strip()),
        "detail": "",
        "warnings": (
            "La base est enregistrée dans un dossier temporaire : elle sera perdue quand le serveur "
            "redémarre ou quand l’hébergement redéploie. Déplacez-la vers un disque persistant "
            "(variable IMAGINE_DB_PATH), ou utilisez la base PostgreSQL de Render (DATABASE_URL)."
            if db_is_ephemeral() else ""
        ),
    }


def pg_sql(sql: str) -> str:
    """Traduit une requête écrite pour SQLite vers le dialecte PostgreSQL."""
    sql = sql.replace("COLLATE NOCASE", "").replace("COLLATE nocase", "")
    sql = sql.replace("BEGIN IMMEDIATE", "BEGIN")
    sql = re.sub(r"\bREAL\b", "DOUBLE PRECISION", sql)
    # Placeholders « ? » -> « %s » (hors littéraux de chaîne).
    out, in_quote = [], False
    for ch in sql:
        if ch == "'":
            in_quote = not in_quote
        if ch == "?" and not in_quote:
            out.append("%s")
        else:
            out.append(ch)
    return "".join(out)


class DBCursor:
    """Curseur minimal commun à SQLite et PostgreSQL."""

    __slots__ = ("_cur",)

    def __init__(self, cur):
        self._cur = cur

    @property
    def rowcount(self) -> int:
        try:
            return int(self._cur.rowcount)
        except Exception:
            return -1

    def fetchone(self):
        return self._cur.fetchone()

    def fetchall(self):
        return self._cur.fetchall()

    def __iter__(self):
        return iter(self._cur)


class DBConnection:
    """Connexion commune : même API que sqlite3 pour le reste du fichier."""

    def __init__(self, con, backend: str):
        self._con = con
        self.backend = backend
        self._last_rowcount = 0

    def execute(self, sql: str, params=()):
        params = tuple(params) if params else ()
        if self.backend == "postgres":
            cur = self._con.cursor()
            cur.execute(pg_sql(sql), params or None)
            wrapper = DBCursor(cur)
        else:
            wrapper = DBCursor(self._con.execute(sql, params))
        self._last_rowcount = wrapper.rowcount
        return wrapper

    def executescript(self, script: str) -> None:
        if self.backend == "postgres":
            for statement in script.split(";"):
                statement = statement.strip()
                if statement:
                    self.execute(statement)
            return
        self._con.executescript(script)

    @property
    def total_changes(self) -> int:
        if self.backend == "postgres":
            return self._last_rowcount
        return self._con.total_changes

    def _pg_transaction_open(self) -> bool:
        # En autocommit, connection.commit()/rollback() sont des no-op : on envoie
        # COMMIT/ROLLBACK au serveur quand une transaction explicite est ouverte.
        status = self._con.info.transaction_status
        return status in (psycopg2.extensions.TRANSACTION_STATUS_INTRANS,
                          psycopg2.extensions.TRANSACTION_STATUS_INERROR)

    def commit(self) -> None:
        if self.backend == "postgres":
            if self._pg_transaction_open():
                self._con.cursor().execute("COMMIT")
            return
        self._con.commit()

    def rollback(self) -> None:
        if self.backend == "postgres":
            if self._pg_transaction_open():
                self._con.cursor().execute("ROLLBACK")
            return
        self._con.rollback()

    def close(self) -> None:
        self._con.close()


def connect() -> DBConnection:
    if USE_POSTGRES:
        if psycopg2 is None:
            raise RuntimeError("Le pilote psycopg2 n’est pas installé : ajoutez psycopg2-binary au requirements.txt.")
        con = psycopg2.connect(
            DATABASE_URL, cursor_factory=RealDictCursor, connect_timeout=15,
            application_name="imagine-events-desk", keepalives=1, keepalives_idle=30,
        )
        con.autocommit = True  # les transactions sont ouvertes explicitement (BEGIN … commit)
        return DBConnection(con, "postgres")
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=20, isolation_level=None)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA busy_timeout=20000")
    return DBConnection(con, "sqlite")


SCHEMA_STATEMENTS = [
    """
    CREATE TABLE IF NOT EXISTS events (
      id TEXT PRIMARY KEY,
      name TEXT NOT NULL,
      event_date TEXT NOT NULL DEFAULT '',
      venue TEXT NOT NULL DEFAULT '',
      status TEXT NOT NULL DEFAULT 'active',
      currency TEXT NOT NULL DEFAULT 'TND',
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS participants (
      id TEXT PRIMARY KEY,
      event_id TEXT NOT NULL REFERENCES events(id) ON DELETE CASCADE,
      name TEXT NOT NULL,
      email TEXT NOT NULL DEFAULT '',
      company TEXT NOT NULL DEFAULT '',
      category TEXT NOT NULL DEFAULT 'participant',
      code TEXT NOT NULL,
      checked_in_at TEXT,
      payment_status TEXT NOT NULL DEFAULT 'unpaid',
      payment_amount DOUBLE PRECISION NOT NULL DEFAULT 0,
      payment_method TEXT NOT NULL DEFAULT '',
      payment_reference TEXT NOT NULL DEFAULT '',
      payment_confirmed_at TEXT,
      payment_confirmed_by TEXT NOT NULL DEFAULT '',
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      UNIQUE(event_id, code)
    )
    """,
    "CREATE INDEX IF NOT EXISTS idx_participants_event_name ON participants(event_id, name)",
    "CREATE INDEX IF NOT EXISTS idx_participants_event_updated ON participants(event_id, updated_at)",
    """
    CREATE TABLE IF NOT EXISTS sync_operations (
      op_id TEXT PRIMARY KEY,
      event_id TEXT NOT NULL REFERENCES events(id) ON DELETE CASCADE,
      participant_id TEXT NOT NULL,
      action TEXT NOT NULL,
      result_json TEXT NOT NULL,
      created_at TEXT NOT NULL
    )
    """,
    "CREATE INDEX IF NOT EXISTS idx_sync_operations_event ON sync_operations(event_id, created_at)",
    """
    CREATE TABLE IF NOT EXISTS users (
      id TEXT PRIMARY KEY,
      username TEXT NOT NULL UNIQUE,
      password_salt TEXT NOT NULL,
      password_hash TEXT NOT NULL,
      role TEXT NOT NULL CHECK(role IN ('admin','desk')),
      active INTEGER NOT NULL DEFAULT 1,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS sessions (
      token_hash TEXT PRIMARY KEY,
      user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
      expires_at TEXT NOT NULL,
      created_at TEXT NOT NULL
    )
    """,
    "CREATE INDEX IF NOT EXISTS idx_sessions_expiry ON sessions(expires_at)",
]

# Colonnes ajoutées après la première version de l’application (migration douce).
EXPECTED_COLUMNS = {
    "events": [("currency", "TEXT NOT NULL DEFAULT 'TND'")],
    "participants": [
        ("payment_status", "TEXT NOT NULL DEFAULT 'unpaid'"),
        ("payment_amount", "DOUBLE PRECISION NOT NULL DEFAULT 0"),
        ("payment_method", "TEXT NOT NULL DEFAULT ''"),
        ("payment_reference", "TEXT NOT NULL DEFAULT ''"),
        ("payment_confirmed_at", "TEXT"),
        ("payment_confirmed_by", "TEXT NOT NULL DEFAULT ''"),
    ],
}


def init_db() -> None:
    if not USE_POSTGRES:
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = connect()
    try:
        if con.backend == "sqlite":
            con.execute("PRAGMA journal_mode=WAL")
        for statement in SCHEMA_STATEMENTS:
            con.execute(statement)
        for table, columns in EXPECTED_COLUMNS.items():
            for column, definition in columns:
                ensure_column(con, table, column, definition)
        if con.backend == "postgres":
            # Recherche insensible à la casse des noms d’utilisateur (exigence d’unicité).
            con.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_users_username_lower ON users (lower(username))")
    finally:
        con.close()


def ensure_column(con: DBConnection, table: str, column: str, definition: str) -> None:
    if con.backend == "postgres":
        existing = {
            row["column_name"] for row in con.execute(
                "SELECT column_name FROM information_schema.columns WHERE table_schema='public' AND table_name=?",
                (table,),
            ).fetchall()
        }
    else:
        existing = {row[1] for row in con.execute(f"PRAGMA table_info({table})")}
    if column not in existing:
        con.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def clean_text(value, limit=500) -> str:
    if value is None:
        return ""
    return str(value).strip()[:limit]


def clean_category(value) -> str:
    value = clean_text(value, 50).lower()
    aliases = {
        "intervenant": "speaker", "intervenants": "speaker", "speaker": "speaker", "orateur": "speaker",
        "vip": "vip", "presse": "press", "press": "press", "media": "press", "média": "press",
        "équipe": "team", "equipe": "team", "staff": "team", "organisateur": "team",
        "participant": "participant", "participants": "participant", "invité": "vip", "invite": "vip"
    }
    normalized = aliases.get(value, value)
    return normalized if normalized in CATEGORIES else "participant"


def normalize_username(value) -> str:
    return clean_text(value, 80).strip().lower()


def clean_currency(value) -> str:
    raw = clean_text(value, 8).strip().upper()
    return raw if re.fullmatch(r"[A-Z]{3}", raw) else DEFAULT_CURRENCY


def clean_payment_status(value) -> str:
    value = clean_text(value, 30).lower()
    aliases = {
        "paye": "paid", "payé": "paid", "payee": "paid", "regle": "paid", "réglé": "paid",
        "paid": "paid", "confirme": "paid", "confirmé": "paid", "recu": "paid", "encaissé": "paid",
        "exonere": "exempt", "exonéré": "exempt", "exempt": "exempt", "dispense": "exempt",
        "dispensé": "exempt", "gratuit": "exempt", "offert": "exempt",
        "unpaid": "unpaid", "non paye": "unpaid", "non payé": "unpaid", "impaye": "unpaid",
        "impayé": "unpaid", "en attente": "unpaid", "pending": "unpaid", "": "unpaid",
    }
    normalized = aliases.get(value, value)
    return normalized if normalized in PAYMENT_STATUSES else "unpaid"


def clean_payment_method(value) -> str:
    value = clean_text(value, 30).lower()
    aliases = {
        "especes": "cash", "espèces": "cash", "cash": "cash", "liquide": "cash", "caisse": "cash",
        "carte": "card", "card": "card", "cb": "card", "carte bancaire": "card", "tpe": "card",
        "virement": "transfer", "transfer": "transfer", "virement bancaire": "transfer", "banque": "transfer",
        "cheque": "cheque", "chèque": "cheque", "check": "cheque",
        "mobile": "mobile", "paiement mobile": "mobile", "d17": "mobile", "flouci": "mobile",
        "edinar": "mobile", "e-dinar": "mobile", "en ligne": "mobile", "online": "mobile",
        "autre": "other", "other": "other", "": "",
    }
    normalized = aliases.get(value, value)
    return normalized if normalized in PAYMENT_METHODS else "other"


def clean_amount(value) -> float:
    """Montant positif, 3 décimales maximum (millimes). Accepte « 1 500,500 »."""
    if value is None or value == "":
        return 0.0
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    for suffix in ("TND", "DT", "EUR", "€", "USD", "$"):
        text = text.replace(suffix, "").replace(suffix.lower(), "")
    if text in ("", "-", "."):
        return 0.0
    try:
        amount = float(text)
    except (TypeError, ValueError):
        raise ValueError("Montant de paiement invalide.") from None
    if amount < 0 or amount > MAX_PAYMENT_AMOUNT:
        raise ValueError("Le montant du paiement doit être compris entre 0 et 10 000 000.")
    return round(amount, 3)


def parse_payment_status_text(value) -> str:
    """Lit une cellule de statut de paiement (« Payé », « Non payé », « Exonéré »…)."""
    text = clean_text(value, 60).lower()
    for accent, plain in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("ç", "c")):
        text = text.replace(accent, plain)
    text = " ".join(text.replace("_", " ").replace("-", " ").split())
    if not text:
        return ""
    if any(word in text for word in ("exon", "dispens", "gratuit", "offert", "exempt")):
        return "exempt"
    if text in ("0", "false", "non", "no", "impaye", "unpaid", "en attente", "pending", "a regler", "non regle"):
        return "unpaid"
    if text in ("1", "true", "oui", "yes", "paye", "paid", "regle", "confirme", "recu", "encaiss", "valide"):
        return "paid"
    if text.startswith(("pay", "regl", "encaiss", "confirm")):
        return "paid"
    if text.startswith("non ") or text.startswith("pas "):
        return "unpaid"
    return ""


def payment_fields(data, actor_name: str = "") -> dict:
    """Normalise un paiement reçu du client ou d’un import."""
    data = data if isinstance(data, dict) else {}
    status = clean_payment_status(data.get("status", data.get("paymentStatus", "unpaid")))
    amount = clean_amount(data.get("amount", data.get("paymentAmount"))) if status == "paid" else 0.0
    method = clean_payment_method(data.get("method", data.get("paymentMethod"))) if status == "paid" else ""
    reference = clean_text(data.get("reference", data.get("paymentReference")), 120) if status == "paid" else ""
    return {
        "status": status, "amount": amount, "method": method,
        "reference": reference,
        "confirmed_by": clean_text(actor_name, 80) if status != "unpaid" else "",
    }


def payment_json(row: sqlite3.Row) -> dict:
    status = row["payment_status"] if "payment_status" in row.keys() else "unpaid"
    return {
        "paymentStatus": status if status in PAYMENT_STATUSES else "unpaid",
        "paymentAmount": float(row["payment_amount"] or 0) if "payment_amount" in row.keys() else 0.0,
        "paymentMethod": row["payment_method"] or "" if "payment_method" in row.keys() else "",
        "paymentReference": row["payment_reference"] or "" if "payment_reference" in row.keys() else "",
        "paymentConfirmedAt": row["payment_confirmed_at"] or "" if "payment_confirmed_at" in row.keys() else "",
        "paymentConfirmedBy": row["payment_confirmed_by"] or "" if "payment_confirmed_by" in row.keys() else "",
    }


def validate_account(username, password):
    username = normalize_username(username)
    if not re.fullmatch(r"[a-z0-9._-]{3,40}", username):
        raise ValueError("Le nom d’utilisateur doit contenir 3 à 40 caractères (lettres, chiffres, point, tiret ou souligné).")
    if not isinstance(password, str) or len(password) < 12 or len(password) > 256:
        raise ValueError("Le mot de passe doit contenir entre 12 et 256 caractères.")
    return username


def make_password_hash(password, salt=None):
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PASSWORD_ITERATIONS)
    return salt.hex(), digest.hex()


def password_matches(password, salt_hex, expected_hex):
    try:
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), PASSWORD_ITERATIONS).hex()
        return hmac.compare_digest(actual, expected_hex)
    except (ValueError, TypeError, AttributeError):
        return False


def public_user(row):
    return {"id":row["id"],"username":row["username"],"role":row["role"],"active":bool(row["active"]),"createdAt":row["created_at"]}


def unique_code(con: sqlite3.Connection, event_id: str) -> str:
    for _ in range(20):
        code = "IE-" + secrets.token_hex(3).upper()
        exists = con.execute("SELECT 1 FROM participants WHERE event_id=? AND code=?", (event_id, code)).fetchone()
        if not exists:
            return code
    return "IE-" + uuid.uuid4().hex[:8].upper()


def participant_json(row: sqlite3.Row) -> dict:
    data = {
        "id": row["id"], "eventId": row["event_id"], "name": row["name"],
        "email": row["email"], "company": row["company"], "category": row["category"],
        "code": row["code"], "checkedIn": bool(row["checked_in_at"]),
        "checkedInAt": row["checked_in_at"] or "", "updatedAt": row["updated_at"]
    }
    data.update(payment_json(row))
    return data


def event_json(row: sqlite3.Row) -> dict:
    keys = row.keys()
    return {
        "id": row["id"], "name": row["name"], "date": row["event_date"],
        "venue": row["venue"], "status": row["status"], "createdAt": row["created_at"],
        "updatedAt": row["updated_at"], "total": int(row["total"] or 0),
        "arrived": int(row["arrived"] or 0),
        "currency": (row["currency"] if "currency" in keys else DEFAULT_CURRENCY) or DEFAULT_CURRENCY,
        "paid": int(row["paid"] or 0) if "paid" in keys else 0,
        "collected": round(float(row["collected"] or 0), 3) if "collected" in keys else 0.0,
    }


class Handler(BaseHTTPRequestHandler):
    server_version = "ImagineDesk/0.1"

    def log_message(self, fmt, *args):
        # Logs compacts du serveur; aucune donnée personnelle n'est écrite ici.
        print("%s - %s" % (self.address_string(), fmt % args))

    def send_bytes(self, status: int, body: bytes, content_type: str, extra_headers=None):
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "same-origin")
        if self.path.startswith("/api/"):
            self.send_header("Cache-Control", "no-store")
        for key, value in (extra_headers or {}).items():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(body)

    def json(self, status: int, payload, extra_headers=None):
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        self.send_bytes(status, body, "application/json; charset=utf-8", extra_headers)

    def read_raw(self, max_bytes=MAX_BODY):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            raise ValueError("Longueur de requête invalide.")
        if length <= 0:
            raise ValueError("Requête vide.")
        if length > max_bytes:
            limit_mb = max(1, max_bytes // (1024 * 1024))
            raise OverflowError(f"Requête trop volumineuse pour cette version (limite {limit_mb} Mo).")
        return self.rfile.read(length)

    def read_json(self):
        raw = self.read_raw()
        try:
            result = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ValueError("JSON invalide.")
        if not isinstance(result, dict):
            raise ValueError("Le corps doit être un objet JSON.")
        return result

    def current_user(self):
        cookie = SimpleCookie()
        try:
            cookie.load(self.headers.get("Cookie", ""))
        except Exception:
            return None
        morsel = cookie.get(SESSION_COOKIE)
        if not morsel or not morsel.value:
            return None
        token_hash = hashlib.sha256(morsel.value.encode("utf-8")).hexdigest()
        con = connect()
        try:
            row = con.execute("""SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id
                              WHERE s.token_hash=? AND s.expires_at>? AND u.active=1""", (token_hash, utc_now())).fetchone()
            return public_user(row) if row else None
        finally:
            con.close()

    def new_session(self, user_id):
        token = secrets.token_urlsafe(32)
        now = utc_now()
        expires = (datetime.now(timezone.utc) + timedelta(seconds=SESSION_TTL_SECONDS)).isoformat(timespec="microseconds")
        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        con = connect()
        try:
            con.execute("DELETE FROM sessions WHERE expires_at<=?", (now,))
            con.execute("INSERT INTO sessions(token_hash,user_id,expires_at,created_at) VALUES(?,?,?,?)", (token_hash,user_id,expires,now))
        finally:
            con.close()
        return token

    def auth_cookie(self, token=None, clear=False):
        if clear:
            value = f"{SESSION_COOKIE}=; Path=/; HttpOnly; SameSite=Strict; Max-Age=0"
        else:
            value = f"{SESSION_COOKIE}={token}; Path=/; HttpOnly; SameSite=Strict; Max-Age={SESSION_TTL_SECONDS}"
        forwarded = self.headers.get("X-Forwarded-Proto", "").lower()
        origin = self.headers.get("Origin", "").lower()
        if os.environ.get("COOKIE_SECURE") == "1" or "https" in forwarded or origin.startswith("https://"):
            value += "; Secure"
        return {"Set-Cookie":value}

    def require_user(self):
        user = self.current_user()
        if not user:
            self.json(401, {"error":"Connexion requise. Connectez-vous à nouveau."})
            return None
        return user

    def require_admin(self, user):
        if not user or user.get("role") != "admin":
            self.json(403, {"error":"Cette action est réservée à l’administrateur."})
            return False
        return True

    def auth_status(self):
        con = connect()
        try:
            count = con.execute("SELECT COUNT(*) AS n FROM users").fetchone()["n"]
        finally:
            con.close()
        user = self.current_user()
        return {"setupRequired":count==0,"authenticated":bool(user),"user":user}

    def create_account(self, username, password, role="desk"):
        username = validate_account(username, password)
        if role not in ("admin", "desk"):
            raise ValueError("Rôle invalide.")
        salt, digest = make_password_hash(password)
        user_id = "usr-" + uuid.uuid4().hex
        now = utc_now()
        con = connect()
        try:
            con.execute("INSERT INTO users(id,username,password_salt,password_hash,role,active,created_at,updated_at) VALUES(?,?,?,?,?,1,?,?)", (user_id,username,salt,digest,role,now,now))
            row = con.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
            return public_user(row)
        finally:
            con.close()

    def login_limited(self):
        key = self.client_address[0] if getattr(self, "client_address", None) else "unknown"
        now = time.time()
        attempts = [t for t in LOGIN_FAILURES.get(key, []) if now-t < 900]
        LOGIN_FAILURES[key] = attempts
        return len(attempts) >= 10

    def record_login_failure(self):
        key = self.client_address[0] if getattr(self, "client_address", None) else "unknown"
        now = time.time()
        LOGIN_FAILURES[key] = [t for t in LOGIN_FAILURES.get(key, []) if now-t < 900] + [now]

    def setup_first_admin(self, body):
        provided_token = body.get("setupToken") if isinstance(body.get("setupToken"), str) else ""
        if SETUP_TOKEN and not hmac.compare_digest(provided_token, SETUP_TOKEN):
            self.json(403, {"error":"Code de configuration invalide. Vérifiez IMAGINE_SETUP_TOKEN auprès de la personne qui administre le serveur."})
            return
        username = validate_account(body.get("username"), body.get("password"))
        salt, digest = make_password_hash(body.get("password"))
        user_id = "usr-" + uuid.uuid4().hex
        now = utc_now()
        con = connect()
        try:
            con.execute("BEGIN IMMEDIATE")
            count = con.execute("SELECT COUNT(*) AS n FROM users").fetchone()["n"]
            if count:
                con.rollback()
                self.json(409, {"error":"Le premier compte administrateur existe déjà. Connectez-vous."})
                return
            con.execute("INSERT INTO users(id,username,password_salt,password_hash,role,active,created_at,updated_at) VALUES(?,?,?,?, 'admin',1,?,?)", (user_id,username,salt,digest,now,now))
            con.commit()
            row = con.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
            token = self.new_session(user_id)
            self.json(201, {"user":public_user(row)}, self.auth_cookie(token))
        except Exception:
            con.rollback()
            raise
        finally:
            con.close()

    def login(self, body):
        if self.login_limited():
            self.json(429, {"error":"Trop de tentatives. Réessayez dans quelques minutes."})
            return
        username = normalize_username(body.get("username"))
        password = body.get("password") if isinstance(body.get("password"), str) else ""
        con = connect()
        try:
            row = con.execute("SELECT * FROM users WHERE lower(username)=lower(?)", (username,)).fetchone()
        finally:
            con.close()
        if not row or not row["active"] or not password_matches(password, row["password_salt"], row["password_hash"]):
            self.record_login_failure()
            self.json(401, {"error":"Nom d’utilisateur ou mot de passe incorrect."})
            return
        key = self.client_address[0] if getattr(self, "client_address", None) else "unknown"
        LOGIN_FAILURES.pop(key, None)
        token = self.new_session(row["id"])
        self.json(200, {"user":public_user(row)}, self.auth_cookie(token))

    def logout(self):
        cookie = SimpleCookie()
        try: cookie.load(self.headers.get("Cookie", ""))
        except Exception: pass
        morsel = cookie.get(SESSION_COOKIE)
        if morsel and morsel.value:
            token_hash = hashlib.sha256(morsel.value.encode("utf-8")).hexdigest()
            con = connect()
            try: con.execute("DELETE FROM sessions WHERE token_hash=?", (token_hash,))
            finally: con.close()
        self.json(200, {"ok":True}, self.auth_cookie(clear=True))

    def update_user(self, user_id, body, actor):
        role = body.get("role")
        active = body.get("active")
        con = connect()
        try:
            target = con.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
            if not target:
                self.json(404, {"error":"Compte introuvable."}); return
            new_role = role if role in ("admin", "desk") else target["role"]
            new_active = int(bool(active)) if isinstance(active, bool) else target["active"]
            if target["role"]=="admin" and (new_role!="admin" or not new_active):
                admins = con.execute("SELECT COUNT(*) AS n FROM users WHERE role='admin' AND active=1").fetchone()["n"]
                if admins <= 1:
                    self.json(400, {"error":"Il faut conserver au moins un administrateur actif."}); return
            if target["id"]==actor["id"] and (new_role!="admin" or not new_active):
                self.json(400, {"error":"Vous ne pouvez pas désactiver ou rétrograder votre propre compte."}); return
            con.execute("UPDATE users SET role=?,active=?,updated_at=? WHERE id=?",(new_role,new_active,utc_now(),user_id))
            if new_role != target["role"] or not new_active:
                con.execute("DELETE FROM sessions WHERE user_id=?",(user_id,))
            row=con.execute("SELECT * FROM users WHERE id=?",(user_id,)).fetchone()
            self.json(200,{"user":public_user(row)})
        finally: con.close()

    def reset_user_password(self, user_id, body):
        password=body.get("password") if isinstance(body.get("password"),str) else ""
        if len(password)<12 or len(password)>256:
            raise ValueError("Le mot de passe doit contenir entre 12 et 256 caractères.")
        salt,digest=make_password_hash(password)
        con=connect()
        try:
            row=con.execute("SELECT id FROM users WHERE id=?",(user_id,)).fetchone()
            if not row:
                self.json(404,{"error":"Compte introuvable."});return
            now=utc_now()
            con.execute("UPDATE users SET password_salt=?,password_hash=?,updated_at=? WHERE id=?",(salt,digest,now,user_id))
            con.execute("DELETE FROM sessions WHERE user_id=?",(user_id,))
            self.json(200,{"ok":True})
        finally:con.close()

    def change_own_password(self, user, body):
        current = body.get("currentPassword") if isinstance(body.get("currentPassword"),str) else ""
        password = body.get("newPassword") if isinstance(body.get("newPassword"),str) else ""
        if len(password)<12 or len(password)>256:
            raise ValueError("Le nouveau mot de passe doit contenir entre 12 et 256 caractères.")
        con=connect()
        try:
            row=con.execute("SELECT * FROM users WHERE id=? AND active=1",(user["id"],)).fetchone()
            if not row or not password_matches(current,row["password_salt"],row["password_hash"]):
                self.json(400,{"error":"Le mot de passe actuel est incorrect."});return
            if password_matches(password,row["password_salt"],row["password_hash"]):
                self.json(400,{"error":"Choisissez un mot de passe différent de l’actuel."});return
            salt,digest=make_password_hash(password)
            now=utc_now()
            con.execute("BEGIN IMMEDIATE")
            con.execute("UPDATE users SET password_salt=?,password_hash=?,updated_at=? WHERE id=?",(salt,digest,now,user["id"]))
            con.execute("DELETE FROM sessions WHERE user_id=?",(user["id"],))
            con.commit()
        finally:con.close()
        token=self.new_session(user["id"])
        self.json(200,{"ok":True},self.auth_cookie(token))

    def export_database_backup(self):
        """Télécharge une sauvegarde cohérente, sans sessions actives.

        SQLite : copie binaire via l’API backup. PostgreSQL : fichier SQL (pg_dump
        léger, tables et données), nommé .sql et rechargeable via psql.
        """
        if USE_POSTGRES:
            self.export_postgres_backup()
            return
        temp_fd, temp_name = tempfile.mkstemp(prefix="imagine-events-export-", suffix=".sqlite3")
        os.close(temp_fd)
        source = None
        target = None
        try:
            source = sqlite3.connect(DB_PATH, timeout=20)
            target = sqlite3.connect(temp_name, timeout=20)
            source.backup(target, pages=256, sleep=0.05)
            target.execute("DELETE FROM sessions")
            target.commit()
            try:
                target.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                target.execute("PRAGMA journal_mode=DELETE")
            except sqlite3.DatabaseError:
                pass
            target.close()
            target = None
            filename = "imagine-events-backup-" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%SZ") + ".sqlite3"
            backup_bytes = Path(temp_name).read_bytes()
            if len(backup_bytes) > MAX_BACKUP_FILE:
                self.json(413, {"error":"La base dépasse la limite de sauvegarde de 50 Mo."})
                return
            self.send_bytes(200, backup_bytes, "application/vnd.sqlite3", {
                "Content-Disposition": f'attachment; filename="{filename}"'
            })
        finally:
            if target is not None:
                target.close()
            if source is not None:
                source.close()
            try:
                os.unlink(temp_name)
            except OSError:
                pass

    def export_postgres_backup(self):
        """Export SQL des tables de l’application (aucune session, aucun mot de passe en clair)."""
        con = connect()
        try:
            lines = [
                "-- Sauvegarde Imagine Events Desk",
                f"-- {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
                "-- Réimport : psql \"$DATABASE_URL\" -f ce-fichier.sql",
                "BEGIN;",
                "DELETE FROM sync_operations;",
                "DELETE FROM participants;",
                "DELETE FROM events;",
                "DELETE FROM sessions;",
                "DELETE FROM users;",
            ]
            # Les comptes sont inclus (empreintes PBKDF2, jamais les mots de passe en clair),
            # comme dans la sauvegarde SQLite ; les sessions sont exclues.
            for table in ("events", "users", "participants", "sync_operations"):
                rows = con.execute(f"SELECT * FROM {table}").fetchall()
                for row in rows:
                    columns = list(row.keys())
                    values = []
                    for column in columns:
                        value = row[column]
                        if value is None:
                            values.append("NULL")
                        elif isinstance(value, (int, float)):
                            values.append(repr(value))
                        else:
                            values.append("'" + str(value).replace("'", "''") + "'")
                    lines.append(
                        f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({', '.join(values)});"
                    )
            lines.append("COMMIT;")
            payload = ("\n".join(lines) + "\n").encode("utf-8")
            if len(payload) > MAX_BACKUP_FILE:
                self.json(413, {"error":"La sauvegarde dépasse la limite de 50 Mo."})
                return
            filename = "imagine-events-backup-" + datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%SZ") + ".sql"
            self.send_bytes(200, payload, "application/sql; charset=utf-8", {
                "Content-Disposition": f'attachment; filename="{filename}"'
            })
        finally:
            con.close()

    def import_database_backup(self):
        """Remplace les données d’événements par une sauvegarde SQLite validée.

        Les comptes déjà présents restent actifs; les anciens comptes sont ajoutés
        si leur identifiant/nom d’utilisateur ne fait pas conflit. Les sessions
        du fichier importé ne sont jamais reprises.
        """
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type.lower():
            raise ValueError("Choisissez un fichier SQLite à importer.")
        # Lire le corps avant de refuser : sinon le client reçoit une connexion coupée.
        raw = self.read_raw(MAX_BACKUP_REQUEST)
        if USE_POSTGRES:
            raise ValueError(
                "La restauration de fichier est réservée au stockage SQLite. En PostgreSQL, "
                "restaurez la sauvegarde .sql avec psql ou depuis le tableau de bord Render."
            )
        envelope = ("MIME-Version: 1.0\r\nContent-Type: " + content_type + "\r\n\r\n").encode("latin-1") + raw
        try:
            message = BytesParser(policy=policy.default).parsebytes(envelope)
            part = next((p for p in message.iter_parts() if p.get_filename()), None)
            if part is None:
                raise ValueError("Aucun fichier de base de données reçu.")
            filename = clean_text(Path(part.get_filename()).name, 250).lower()
            if not filename.endswith((".sqlite3", ".sqlite", ".db")):
                raise ValueError("Utilisez une sauvegarde .sqlite3, .sqlite ou .db.")
            file_bytes = part.get_payload(decode=True)
            if not file_bytes:
                raise ValueError("Le fichier SQLite est vide.")
            if len(file_bytes) > MAX_BACKUP_FILE:
                raise OverflowError("La sauvegarde dépasse la limite de 50 Mo.")
            if not file_bytes.startswith(b"SQLite format 3\x00"):
                raise ValueError("Le fichier ne semble pas être une base SQLite valide.")
        except OverflowError:
            raise
        except ValueError:
            raise
        except Exception:
            raise ValueError("Impossible de lire le formulaire d’import. Choisissez un fichier SQLite valide.")

        temp_fd, temp_name = tempfile.mkstemp(prefix="imagine-events-import-", suffix=".sqlite3")
        os.close(temp_fd)
        source = None
        con = None
        try:
            Path(temp_name).write_bytes(file_bytes)
            source = sqlite3.connect(f"file:{temp_name}?mode=ro&immutable=1", uri=True, timeout=20)
            source.row_factory = sqlite3.Row
            check = source.execute("PRAGMA quick_check").fetchone()
            if not check or check[0] != "ok":
                raise ValueError("La sauvegarde est endommagée; aucun changement n’a été appliqué.")

            tables = {row[0] for row in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if not {"events", "participants"}.issubset(tables):
                raise ValueError("Cette base ne contient pas les tables événements et participants attendues.")

            def columns(table):
                return {row[1] for row in source.execute(f"PRAGMA table_info({table})")}

            event_columns = columns("events")
            participant_columns = columns("participants")
            if not {"id", "name"}.issubset(event_columns):
                raise ValueError("Schéma d’événements non compatible.")
            if not {"id", "event_id", "name"}.issubset(participant_columns):
                raise ValueError("Schéma de participants non compatible.")
            sync_columns = columns("sync_operations") if "sync_operations" in tables else set()
            required_sync = {"op_id", "event_id", "participant_id", "action", "result_json", "created_at"}
            if sync_columns and not required_sync.issubset(sync_columns):
                raise ValueError("Schéma d’historique de synchronisation non compatible.")
            user_columns = columns("users") if "users" in tables else set()
            required_users = {"id", "username", "password_salt", "password_hash", "role", "active", "created_at", "updated_at"}
            if user_columns and not required_users.issubset(user_columns):
                raise ValueError("Schéma de comptes non compatible.")

            con = connect()
            con.execute("BEGIN IMMEDIATE")
            con.execute("DELETE FROM sync_operations")
            con.execute("DELETE FROM participants")
            con.execute("DELETE FROM events")
            event_count = participant_count = operation_count = imported_users = skipped_users = 0
            event_ids = set()

            for raw_event in source.execute("SELECT * FROM events"):
                item = dict(raw_event)
                event_id = clean_text(item.get("id"), 100)
                name = clean_text(item.get("name"), 180)
                if not event_id or not name:
                    raise ValueError("La sauvegarde contient un événement sans identifiant ou sans nom.")
                now = utc_now()
                status = clean_text(item.get("status", "active"), 20)
                if status not in ("active", "archived"):
                    status = "active"
                con.execute("INSERT INTO events(id,name,event_date,venue,status,currency,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", (
                    event_id, name, clean_text(item.get("event_date", item.get("date", "")), 20),
                    clean_text(item.get("venue", ""), 180), status, clean_currency(item.get("currency")),
                    clean_text(item.get("created_at", ""), 64) or now,
                    clean_text(item.get("updated_at", ""), 64) or now
                ))
                event_ids.add(event_id)
                event_count += 1

            for raw_person in source.execute("SELECT * FROM participants"):
                item = dict(raw_person)
                event_id = clean_text(item.get("event_id"), 100)
                person_id = clean_text(item.get("id"), 100)
                name = clean_text(item.get("name"), 180)
                if not event_id or event_id not in event_ids or not person_id or not name:
                    raise ValueError("La sauvegarde contient un participant incomplet ou lié à un congrès absent.")
                now = utc_now()
                code = clean_text(item.get("code", ""), 40)
                if not code or con.execute("SELECT 1 FROM participants WHERE event_id=? AND lower(code)=lower(?)", (event_id, code)).fetchone():
                    code = unique_code(con, event_id)
                checked = clean_text(item.get("checked_in_at", ""), 64) or None
                pay_status = clean_payment_status(item.get("payment_status") or "unpaid")
                try:
                    pay_amount = clean_amount(item.get("payment_amount") or 0) if pay_status == "paid" else 0.0
                except ValueError:
                    pay_amount = 0.0
                pay_method = clean_payment_method(item.get("payment_method") or "") if pay_status == "paid" else ""
                pay_reference = clean_text(item.get("payment_reference") or "", 120) if pay_status == "paid" else ""
                pay_at = clean_text(item.get("payment_confirmed_at") or "", 64) if pay_status != "unpaid" else None
                pay_by = clean_text(item.get("payment_confirmed_by") or "", 80) if pay_status != "unpaid" else ""
                con.execute("""INSERT INTO participants(
                    id,event_id,name,email,company,category,code,checked_in_at,
                    payment_status,payment_amount,payment_method,payment_reference,payment_confirmed_at,payment_confirmed_by,
                    created_at,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                    person_id, event_id, name, clean_text(item.get("email", ""), 180),
                    clean_text(item.get("company", ""), 180), clean_category(item.get("category", "participant")),
                    code, checked,
                    pay_status, pay_amount, pay_method, pay_reference, pay_at, pay_by,
                    clean_text(item.get("created_at", ""), 64) or now,
                    now  # force les postes déjà ouverts à recharger la liste restaurée
                ))
                participant_count += 1

            if sync_columns:
                for raw_operation in source.execute("SELECT op_id,event_id,participant_id,action,result_json,created_at FROM sync_operations"):
                    item = dict(raw_operation)
                    op_id = clean_text(item.get("op_id"), 100)
                    event_id = clean_text(item.get("event_id"), 100)
                    action = clean_text(item.get("action"), 30)
                    if not op_id or event_id not in event_ids or action not in ("create", "checkin", "undo", "payment"):
                        raise ValueError("La sauvegarde contient un historique de synchronisation invalide.")
                    con.execute("INSERT INTO sync_operations(op_id,event_id,participant_id,action,result_json,created_at) VALUES(?,?,?,?,?,?)", (
                        op_id, event_id, clean_text(item.get("participant_id"), 100), action,
                        clean_text(item.get("result_json", "{}"), 4000) or "{}",
                        clean_text(item.get("created_at", ""), 64) or utc_now()
                    ))
                    operation_count += 1

            if user_columns:
                for raw_user in source.execute("SELECT id,username,password_salt,password_hash,role,active,created_at,updated_at FROM users"):
                    item = dict(raw_user)
                    username = normalize_username(item.get("username"))
                    salt = clean_text(item.get("password_salt"), 128)
                    password_hash = clean_text(item.get("password_hash"), 128)
                    role = clean_text(item.get("role"), 20)
                    try:
                        salt_bytes, hash_bytes = bytes.fromhex(salt), bytes.fromhex(password_hash)
                    except ValueError:
                        skipped_users += 1
                        continue
                    if (not re.fullmatch(r"[a-z0-9._-]{3,40}", username) or role not in ("admin", "desk")
                            or len(salt_bytes) != 16 or len(hash_bytes) != 32):
                        skipped_users += 1
                        continue
                    before = con.total_changes
                    values = (
                        clean_text(item.get("id"), 100) or "usr-" + uuid.uuid4().hex,
                        username, salt, password_hash, role, int(bool(item.get("active", 1))),
                        clean_text(item.get("created_at", ""), 64) or utc_now(),
                        clean_text(item.get("updated_at", ""), 64) or utc_now()
                    )
                    if con.backend == "postgres":
                        # Pas de « INSERT OR IGNORE » en PostgreSQL : on isole la ligne
                        # dans un point de sauvegarde et on ignore un conflit d’identifiant.
                        con.execute("SAVEPOINT user_import")
                        try:
                            con.execute("INSERT INTO users(id,username,password_salt,password_hash,role,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", values)
                            con.execute("RELEASE SAVEPOINT user_import")
                            imported_users += 1
                        except DB_INTEGRITY_ERRORS:
                            con.execute("ROLLBACK TO SAVEPOINT user_import")
                            skipped_users += 1
                    else:
                        con.execute("INSERT OR IGNORE INTO users(id,username,password_salt,password_hash,role,active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", values)
                        if con.total_changes > before:
                            imported_users += 1
                        else:
                            skipped_users += 1

            con.commit()
            self.json(200, {
                "ok": True, "events": event_count, "participants": participant_count,
                "syncOperations": operation_count, "accountsAdded": imported_users,
                "accountsSkipped": skipped_users,
                "message": "Les événements et participants ont été remplacés. Les comptes actuels sont conservés; les sessions de la sauvegarde ne sont pas importées."
            })
        except sqlite3.DatabaseError as exc:
            if con is not None:
                con.rollback()
            raise ValueError("La sauvegarde SQLite est illisible/incompatible ou la base est occupée; aucun changement n’a été appliqué.") from exc
        except Exception:
            if con is not None:
                con.rollback()
            raise
        finally:
            if source is not None:
                source.close()
            if con is not None:
                con.close()
            try:
                os.unlink(temp_name)
            except OSError:
                pass

    def do_GET(self):
        parsed = urlparse(self.path)
        path = unquote(parsed.path)
        if path == "/api/auth/status":
            self.json(200, self.auth_status())
            return
        if path == "/api/health":
            self.json(200, {
                "ok": True, "time": utc_now(), "storage": db_backend(),
                "database": db_display_path(), "persistent": not db_is_ephemeral(),
            })
            return
        user = None
        if path.startswith("/api/"):
            user = self.require_user()
            if not user:
                return
            if path == "/api/users":
                if not self.require_admin(user):
                    return
                con = connect()
                try:
                    rows = con.execute("SELECT * FROM users ORDER BY username COLLATE NOCASE").fetchall()
                    self.json(200, {"users":[public_user(row) for row in rows]})
                finally:
                    con.close()
                return
            if path == "/api/admin/storage":
                if not self.require_admin(user):
                    return
                self.json(200, storage_info())
                return
            if path == "/api/admin/backup/export":
                if not self.require_admin(user):
                    return
                self.export_database_backup()
                return
        if path == "/api/events":
            con = connect()
            try:
                rows = con.execute("""
                  SELECT e.*, COUNT(p.id) AS total,
                    SUM(CASE WHEN p.checked_in_at IS NOT NULL THEN 1 ELSE 0 END) AS arrived,
                    SUM(CASE WHEN p.payment_status='paid' THEN 1 ELSE 0 END) AS paid,
                    SUM(CASE WHEN p.payment_status='paid' THEN p.payment_amount ELSE 0 END) AS collected
                  FROM events e LEFT JOIN participants p ON p.event_id=e.id
                  GROUP BY e.id ORDER BY e.created_at DESC
                """).fetchall()
                self.json(200, {"events": [event_json(row) for row in rows], "serverTime": utc_now()})
            finally:
                con.close()
            return
        match = re.fullmatch(r"/api/events/([^/]+)/participants", path)
        if match:
            event_id = match.group(1)
            query = parse_qs(parsed.query)
            since = query.get("since", [""])[0]
            con = connect()
            try:
                exists = con.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone()
                if not exists:
                    self.json(404, {"error": "Événement introuvable."})
                    return
                if since:
                    rows = con.execute("SELECT * FROM participants WHERE event_id=? AND updated_at>? ORDER BY name COLLATE NOCASE", (event_id, since)).fetchall()
                else:
                    rows = con.execute("SELECT * FROM participants WHERE event_id=? ORDER BY name COLLATE NOCASE", (event_id,)).fetchall()
                self.json(200, {"participants": [participant_json(row) for row in rows], "serverTime": utc_now()})
            finally:
                con.close()
            return
        if path == "/" or path == "/index.html":
            self.serve_static("index.html")
            return
        if path.startswith("/api/"):
            self.json(404, {"error": "Route API introuvable."})
            return
        self.serve_static(path.lstrip("/"))

    def serve_static(self, relative: str):
        safe = (PUBLIC / relative).resolve()
        try:
            safe.relative_to(PUBLIC.resolve())
        except ValueError:
            self.json(403, {"error": "Accès refusé."})
            return
        if not safe.is_file():
            self.json(404, {"error": "Fichier introuvable."})
            return
        suffix = safe.suffix.lower()
        mime = {".html":"text/html; charset=utf-8", ".css":"text/css; charset=utf-8", ".js":"application/javascript; charset=utf-8", ".svg":"image/svg+xml", ".webmanifest":"application/manifest+json; charset=utf-8", ".json":"application/json; charset=utf-8", ".xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(suffix, "application/octet-stream")
        self.send_bytes(200, safe.read_bytes(), mime, {"Cache-Control":"no-cache"})

    def do_POST(self):
        path = unquote(urlparse(self.path).path)
        try:
            if path == "/api/auth/setup":
                body = self.read_json()
                self.setup_first_admin(body)
                return
            if path == "/api/auth/login":
                body = self.read_json()
                self.login(body)
                return
            if path == "/api/auth/logout":
                self.logout()
                return

            user = self.require_user()
            if not user:
                return
            if path == "/api/auth/password":
                self.change_own_password(user, self.read_json())
                return
            if path == "/api/admin/backup/import":
                if not self.require_admin(user):
                    return
                self.import_database_backup()
                return
            if path == "/api/users":
                if not self.require_admin(user): return
                body = self.read_json()
                account = self.create_account(body.get("username"), body.get("password"), body.get("role", "desk"))
                self.json(201, {"user":account})
                return
            user_password = re.fullmatch(r"/api/users/([^/]+)/password", path)
            if user_password:
                if not self.require_admin(user): return
                self.reset_user_password(user_password.group(1), self.read_json())
                return

            excel_match = re.fullmatch(r"/api/events/([^/]+)/participants/import-excel", path)
            if excel_match:
                if not self.require_admin(user): return
                self.import_excel(excel_match.group(1), user)
                return
            body = self.read_json()
            if path == "/api/events":
                if not self.require_admin(user): return
                self.create_event(body)
                return
            match = re.fullmatch(r"/api/events/([^/]+)/participants/import", path)
            if match:
                if not self.require_admin(user): return
                self.import_people(match.group(1), body, user)
                return
            match = re.fullmatch(r"/api/events/([^/]+)/participants", path)
            if match:
                self.create_person(match.group(1), body, user)
                return
            match = re.fullmatch(r"/api/events/([^/]+)/sync", path)
            if match:
                self.sync_operations(match.group(1), body, user)
                return
            self.json(404, {"error":"Route API introuvable."})
        except OverflowError as exc:
            self.json(413, {"error":str(exc)})
        except ValueError as exc:
            self.json(400, {"error":str(exc)})
        except DB_INTEGRITY_ERRORS as exc:
            print("POST integrity error:", repr(exc)[:300], flush=True)
            self.json(409, {"error":"Conflit de données. Vérifiez le nom d’utilisateur, le code badge ou la liste."})
        except Exception as exc:
            print("POST error:", repr(exc), flush=True)
            self.json(500, {"error":"Erreur interne du serveur."})

    def do_PATCH(self):
        path = unquote(urlparse(self.path).path)
        try:
            user = self.require_user()
            if not user:
                return
            users_match = re.fullmatch(r"/api/users/([^/]+)", path)
            if users_match:
                if not self.require_admin(user): return
                self.update_user(users_match.group(1), self.read_json(), user)
                return
            if not self.require_admin(user):
                return
            body = self.read_json()
            match = re.fullmatch(r"/api/events/([^/]+)", path)
            if match:
                self.update_event(match.group(1), body)
                return
            match = re.fullmatch(r"/api/events/([^/]+)/participants/([^/]+)", path)
            if match:
                self.update_person(match.group(1), match.group(2), body)
                return
            self.json(404, {"error":"Route API introuvable."})
        except OverflowError as exc:
            self.json(413, {"error":str(exc)})
        except ValueError as exc:
            self.json(400, {"error":str(exc)})
        except DB_INTEGRITY_ERRORS:
            self.json(409, {"error":"Conflit de données."})
        except Exception as exc:
            print("PATCH error:", repr(exc), flush=True)
            self.json(500, {"error":"Erreur interne du serveur."})

    def create_event(self, body):
        name = clean_text(body.get("name"), 180)
        if not name:
            raise ValueError("Le nom de l’événement est obligatoire.")
        event_id = "evt-" + uuid.uuid4().hex
        now = utc_now()
        date = clean_text(body.get("date"), 20)
        venue = clean_text(body.get("venue"), 180)
        currency = clean_currency(body.get("currency"))
        con = connect()
        try:
            con.execute("INSERT INTO events(id,name,event_date,venue,status,currency,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)", (
                event_id, name, date, venue, "active", currency, now, now
            ))
            self.json(201, {"event":{"id":event_id,"name":name,"date":date,"venue":venue,"status":"active","currency":currency,"createdAt":now,"updatedAt":now,"total":0,"arrived":0,"paid":0,"collected":0}})
        finally:
            con.close()

    def update_event(self, event_id, body):
        con = connect()
        try:
            row = con.execute("SELECT * FROM events WHERE id=?", (event_id,)).fetchone()
            if not row:
                self.json(404, {"error":"Événement introuvable."})
                return
            name = clean_text(body.get("name", row["name"]), 180) or row["name"]
            date = clean_text(body.get("date", row["event_date"]), 20)
            venue = clean_text(body.get("venue", row["venue"]), 180)
            status = clean_text(body.get("status", row["status"]), 20)
            currency = clean_currency(body.get("currency", row["currency"]))
            if status not in ("active", "archived"):
                raise ValueError("Statut d’événement invalide.")
            now = utc_now()
            con.execute("UPDATE events SET name=?,event_date=?,venue=?,status=?,currency=?,updated_at=? WHERE id=?", (name,date,venue,status,currency,now,event_id))
            self.json(200, {"ok":True})
        finally:
            con.close()

    def create_person(self, event_id, body, user=None):
        name = clean_text(body.get("name"), 180)
        if not name:
            raise ValueError("Le nom du participant est obligatoire.")
        con = connect()
        try:
            if not con.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
                self.json(404, {"error":"Événement introuvable."})
                return
            person_id = clean_text(body.get("id"), 100) or "p-" + uuid.uuid4().hex
            now = utc_now()
            code = clean_text(body.get("code"), 40) or unique_code(con, event_id)
            checked = now if body.get("checkedIn") else None
            payment = payment_fields(body, user["username"] if user else "")
            confirmed_at = now if payment["status"] != "unpaid" else None
            con.execute("""INSERT INTO participants(
                id,event_id,name,email,company,category,code,checked_in_at,
                payment_status,payment_amount,payment_method,payment_reference,payment_confirmed_at,payment_confirmed_by,
                created_at,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                person_id,event_id,name,clean_text(body.get("email"),180),clean_text(body.get("company"),180),
                clean_category(body.get("category")),code,checked,
                payment["status"],payment["amount"],payment["method"],payment["reference"],confirmed_at,payment["confirmed_by"],
                now,now
            ))
            row = con.execute("SELECT * FROM participants WHERE id=?", (person_id,)).fetchone()
            self.json(201, {"participant":participant_json(row)})
        finally:
            con.close()

    def import_excel(self, event_id, user=None):
        """Lit la première feuille d'un classeur .xlsx puis réutilise l'import commun."""
        content_type = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in content_type.lower():
            raise ValueError("Envoyez le classeur Excel au format multipart/form-data.")
        raw = self.read_raw()
        try:
            envelope = ("MIME-Version: 1.0\r\nContent-Type: " + content_type + "\r\n\r\n").encode("latin-1") + raw
            message = BytesParser(policy=policy.default).parsebytes(envelope)
            part = next((p for p in message.iter_parts() if p.get_filename()), None)
            if part is None:
                raise ValueError("Aucun fichier Excel n’a été reçu.")
            filename = clean_text(part.get_filename(), 250).lower()
            if not filename.endswith(".xlsx"):
                raise ValueError("Utilisez un classeur Excel récent au format .xlsx (pas .xls).")
            file_bytes = part.get_payload(decode=True)
            if not file_bytes:
                raise ValueError("Le fichier Excel est vide.")
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                workbook.close()
                raise ValueError("La première feuille Excel ne contient pas de ligne d’en-tête.")

            def cell_text(value):
                if value is None:
                    return ""
                if hasattr(value, "isoformat"):
                    return value.isoformat()
                if isinstance(value, float) and value.is_integer():
                    return str(int(value))
                return str(value).strip()

            def norm_header(value):
                value = unicodedata.normalize("NFKD", cell_text(value).lower().strip())
                return " ".join("".join(ch for ch in value if not unicodedata.combining(ch)).replace("_", " ").replace("-", " ").split())

            headers = [cell_text(value) for value in header_row]
            normalized = [norm_header(value) for value in headers]
            aliases = {
                "full": ["nom complet", "full name", "name", "participant", "nom et prenom", "participant name"],
                "first": ["prenom", "first name", "firstname", "given name"],
                "last": ["nom", "surname", "lastname", "last name", "family name"],
                "email": ["email", "e mail", "courriel", "adresse email", "mail"],
                "company": ["organisation", "organization", "entreprise", "company", "societe", "institution", "structure"],
                "category": ["categorie", "category", "type", "profil", "acces", "access"],
                "code": ["code badge", "code", "reference", "ref", "badge", "ticket", "identifiant"],
                "status": ["statut", "status", "presence", "arrivee", "check in", "checkin", "checked in"],
                "payment": ["paiement", "payment", "statut paiement", "statut du paiement", "reglement", "regle", "paye", "payment status", "paid"],
                "amount": ["montant", "montant paye", "montant du paiement", "amount", "prix", "tarif", "montant tnd"],
                "method": ["moyen de paiement", "mode de paiement", "moyen", "payment method", "mode de reglement"]
            }
            columns = {key: next((i for i, value in enumerate(normalized) if value in candidates), -1) for key, candidates in aliases.items()}
            if columns["full"] < 0 and columns["first"] < 0 and columns["last"] < 0:
                workbook.close()
                raise ValueError("Colonne de nom introuvable. Utilisez « Nom complet » ou « Prénom » et « Nom ».")

            def value_at(row, key):
                index = columns[key]
                return cell_text(row[index]) if index >= 0 and index < len(row) else ""

            participants = []
            for line_number, row in enumerate(iterator, start=2):
                if line_number > 10001:
                    workbook.close()
                    raise ValueError("Limite de 10 000 participants par import Excel.")
                full_name = value_at(row, "full")
                name = full_name or " ".join(part for part in (value_at(row, "first"), value_at(row, "last")) if part)
                if not name:
                    continue
                category_text = norm_header(value_at(row, "category"))
                if "vip" in category_text or "invite" in category_text:
                    category = "vip"
                elif any(word in category_text for word in ("interven", "speaker", "orateur", "conferenc")):
                    category = "speaker"
                elif any(word in category_text for word in ("presse", "press", "media")):
                    category = "press"
                elif any(word in category_text for word in ("equipe", "staff", "organisateur")):
                    category = "team"
                else:
                    category = "participant"
                status_text = norm_header(value_at(row, "status"))
                checked_in = status_text in ("1", "true", "oui", "yes", "present", "arrive", "arrivee", "checked in", "checked-in") or status_text.startswith("arriv") or status_text.startswith("present")
                entry = {
                    "id": "p-" + uuid.uuid4().hex,
                    "name": name[:180],
                    "email": value_at(row, "email")[:180],
                    "company": value_at(row, "company")[:180],
                    "category": category,
                    "code": value_at(row, "code")[:40],
                    "checkedIn": bool(checked_in)
                }
                payment_status = parse_payment_status_text(value_at(row, "payment"))
                try:
                    payment_amount = clean_amount(value_at(row, "amount"))
                except ValueError:
                    payment_amount = 0.0
                if payment_status:
                    entry["paymentStatus"] = payment_status
                if payment_amount > 0:
                    entry["paymentAmount"] = payment_amount
                    entry.setdefault("paymentStatus", "paid")
                method = clean_payment_method(value_at(row, "method"))
                if method and entry.get("paymentStatus") == "paid":
                    entry["paymentMethod"] = method
                participants.append(entry)
            workbook.close()
            if not participants:
                raise ValueError("Aucun participant avec un nom valide n’a été trouvé dans la feuille Excel.")
        except ImportError:
            raise ValueError("La lecture Excel n’est pas installée sur ce serveur. Installez les dépendances du fichier requirements.txt.")
        except ValueError:
            raise
        except Exception as exc:
            print("Excel import parse error:", repr(exc))
            raise ValueError("Impossible de lire ce classeur. Vérifiez qu’il s’agit d’un fichier .xlsx valide.")
        self.import_people(event_id, {"participants": participants}, user)

    def import_people(self, event_id, body, user=None):
        people = body.get("participants")
        if not isinstance(people, list) or not people:
            raise ValueError("La liste de participants est vide.")
        if len(people) > 10000:
            raise ValueError("Limite de 10 000 lignes par import.")
        con = connect()
        added = updated = 0
        changed_ids = []
        try:
            if not con.execute("SELECT 1 FROM events WHERE id=?", (event_id,)).fetchone():
                self.json(404, {"error":"Événement introuvable."})
                return
            con.execute("BEGIN IMMEDIATE")
            for raw in people:
                if not isinstance(raw, dict):
                    continue
                name = clean_text(raw.get("name"), 180)
                if not name:
                    continue
                email = clean_text(raw.get("email"), 180)
                company = clean_text(raw.get("company"), 180)
                category = clean_category(raw.get("category"))
                code = clean_text(raw.get("code"), 40)
                existing = None
                if email:
                    existing = con.execute("SELECT * FROM participants WHERE event_id=? AND lower(email)=lower(?) LIMIT 1", (event_id,email)).fetchone()
                if not existing and code:
                    existing = con.execute("SELECT * FROM participants WHERE event_id=? AND lower(code)=lower(?) LIMIT 1", (event_id,code)).fetchone()
                now = utc_now()
                # Un paiement déjà confirmé n’est écrasé que si la ligne d’import le précise.
                payment_provided = any(key in raw for key in ("paymentStatus", "paymentAmount", "paymentMethod", "paymentReference"))
                payment = payment_fields(raw, user["username"] if user else "") if payment_provided else None
                if existing:
                    new_code = code or existing["code"]
                    # Keep the saved arrival state unless the import explicitly sends it.
                    check_value = raw.get("checkedIn")
                    checked_at = existing["checked_in_at"]
                    if check_value is True and not checked_at: checked_at = now
                    if check_value is False: checked_at = None
                    if payment is not None:
                        payment_at = now if payment["status"] != "unpaid" else None
                        con.execute("""UPDATE participants SET name=?,email=?,company=?,category=?,code=?,checked_in_at=?,
                            payment_status=?,payment_amount=?,payment_method=?,payment_reference=?,payment_confirmed_at=?,payment_confirmed_by=?,
                            updated_at=? WHERE id=?""", (
                            name,email or existing["email"],company,category,new_code,checked_at,
                            payment["status"],payment["amount"],payment["method"],payment["reference"],payment_at,payment["confirmed_by"],
                            now,existing["id"]
                        ))
                    else:
                        con.execute("UPDATE participants SET name=?,email=?,company=?,category=?,code=?,checked_in_at=?,updated_at=? WHERE id=?", (
                            name,email or existing["email"],company,category,new_code,checked_at,now,existing["id"]
                        ))
                    changed_ids.append(existing["id"]); updated += 1
                else:
                    person_id = clean_text(raw.get("id"), 100) or "p-" + uuid.uuid4().hex
                    if not code or con.execute("SELECT 1 FROM participants WHERE event_id=? AND code=?", (event_id,code)).fetchone():
                        code = unique_code(con,event_id)
                    checked_at = now if raw.get("checkedIn") is True else None
                    payment = payment or payment_fields({}, user["username"] if user else "")
                    payment_at = now if payment["status"] != "unpaid" else None
                    con.execute("""INSERT INTO participants(
                        id,event_id,name,email,company,category,code,checked_in_at,
                        payment_status,payment_amount,payment_method,payment_reference,payment_confirmed_at,payment_confirmed_by,
                        created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                        person_id,event_id,name,email,company,category,code,checked_at,
                        payment["status"],payment["amount"],payment["method"],payment["reference"],payment_at,payment["confirmed_by"],
                        now,now
                    ))
                    changed_ids.append(person_id); added += 1
            con.commit()
            rows=[]
            if changed_ids:
                placeholders=','.join('?' for _ in changed_ids)
                rows=con.execute(f"SELECT * FROM participants WHERE id IN ({placeholders}) ORDER BY name COLLATE NOCASE", changed_ids).fetchall()
            self.json(200,{"added":added,"updated":updated,"participants":[participant_json(row) for row in rows],"serverTime":utc_now()})
        except Exception:
            con.rollback()
            raise
        finally:
            con.close()

    def update_person(self, event_id, person_id, body):
        con=connect()
        try:
            row=con.execute("SELECT * FROM participants WHERE id=? AND event_id=?",(person_id,event_id)).fetchone()
            if not row:
                self.json(404,{"error":"Participant introuvable."}); return
            name=clean_text(body.get("name",row["name"]),180) or row["name"]
            email=clean_text(body.get("email",row["email"]),180)
            company=clean_text(body.get("company",row["company"]),180)
            category=clean_category(body.get("category",row["category"]))
            code=clean_text(body.get("code",row["code"]),40) or row["code"]
            now=utc_now()
            con.execute("UPDATE participants SET name=?,email=?,company=?,category=?,code=?,updated_at=? WHERE id=?",(name,email,company,category,code,now,person_id))
            fresh=con.execute("SELECT * FROM participants WHERE id=?",(person_id,)).fetchone()
            self.json(200,{"participant":participant_json(fresh)})
        finally: con.close()

    def sync_operations(self, event_id, body, user=None):
        operations=body.get("operations",[])
        if not isinstance(operations,list): raise ValueError("La file de synchronisation est invalide.")
        if len(operations)>500: raise ValueError("Synchronisez au maximum 500 actions à la fois.")
        con=connect()
        try:
            if not con.execute("SELECT 1 FROM events WHERE id=?",(event_id,)).fetchone():
                self.json(404,{"error":"Événement introuvable."}); return
            con.execute("BEGIN IMMEDIATE")
            results=[]; changed_ids=[]
            for op in operations:
                if not isinstance(op,dict): continue
                op_id=clean_text(op.get("opId"),100)
                action=clean_text(op.get("action"),30)
                person_id=clean_text(op.get("participantId"),100)
                if not op_id or not person_id or action not in ("create","checkin","undo","payment"):
                    results.append({"opId":op_id,"ok":False,"error":"Action invalide."}); continue
                previous=con.execute("SELECT result_json FROM sync_operations WHERE op_id=?",(op_id,)).fetchone()
                if previous:
                    try: prior=json.loads(previous["result_json"])
                    except Exception: prior={"opId":op_id,"ok":True}
                    results.append(prior); continue
                now=utc_now()
                if action=="create":
                    raw=op.get("participant") or {}
                    name=clean_text(raw.get("name"),180)
                    if not name:
                        result={"opId":op_id,"ok":False,"error":"Nom manquant."}; results.append(result); continue
                    existing=con.execute("SELECT * FROM participants WHERE id=? AND event_id=?",(person_id,event_id)).fetchone()
                    if not existing:
                        code=clean_text(raw.get("code"),40)
                        if not code or con.execute("SELECT 1 FROM participants WHERE event_id=? AND code=?",(event_id,code)).fetchone(): code=unique_code(con,event_id)
                        con.execute("""INSERT INTO participants(
                            id,event_id,name,email,company,category,code,checked_in_at,
                            payment_status,payment_amount,payment_method,payment_reference,payment_confirmed_at,payment_confirmed_by,
                            created_at,updated_at
                        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(
                            person_id,event_id,name,clean_text(raw.get("email"),180),clean_text(raw.get("company"),180),
                            clean_category(raw.get("category")),code,None,
                            "unpaid",0,"","",None,"",now,now
                        ))
                    changed_ids.append(person_id)
                    result={"opId":op_id,"ok":True,"action":"create","participantId":person_id}
                elif action=="payment":
                    person=con.execute("SELECT * FROM participants WHERE id=? AND event_id=?",(person_id,event_id)).fetchone()
                    if not person:
                        result={"opId":op_id,"ok":False,"error":"Participant non présent sur le serveur.","participantId":person_id}; results.append(result); continue
                    try:
                        payment=payment_fields(op.get("payment") or {}, user["username"] if user else "")
                    except ValueError as exc:
                        result={"opId":op_id,"ok":False,"error":str(exc),"participantId":person_id}; results.append(result); continue
                    confirmed_at=now if payment["status"]!="unpaid" else None
                    con.execute("""UPDATE participants SET payment_status=?,payment_amount=?,payment_method=?,payment_reference=?,
                        payment_confirmed_at=?,payment_confirmed_by=?,updated_at=? WHERE id=?""",(
                        payment["status"],payment["amount"],payment["method"],payment["reference"],
                        confirmed_at,payment["confirmed_by"],now,person_id
                    ))
                    changed_ids.append(person_id)
                    result={"opId":op_id,"ok":True,"action":"payment","participantId":person_id,
                            "paymentStatus":payment["status"],"paymentAmount":payment["amount"],
                            "paymentConfirmedAt":confirmed_at or ""}
                else:
                    person=con.execute("SELECT * FROM participants WHERE id=? AND event_id=?",(person_id,event_id)).fetchone()
                    if not person:
                        result={"opId":op_id,"ok":False,"error":"Participant non présent sur le serveur.","participantId":person_id}; results.append(result); continue
                    if action=="checkin":
                        checked=person["checked_in_at"] or now
                    else:
                        checked=None
                    con.execute("UPDATE participants SET checked_in_at=?,updated_at=? WHERE id=?",(checked,now,person_id))
                    changed_ids.append(person_id)
                    result={"opId":op_id,"ok":True,"action":action,"participantId":person_id,"checkedInAt":checked or ""}
                result_json=json.dumps(result,ensure_ascii=False,separators=(",",":"))
                con.execute("INSERT INTO sync_operations(op_id,event_id,participant_id,action,result_json,created_at) VALUES(?,?,?,?,?,?)",(op_id,event_id,person_id,action,result_json,now))
                results.append(result)
            con.commit()
            fresh=[]
            if changed_ids:
                unique_ids=list(dict.fromkeys(changed_ids)); placeholders=','.join('?' for _ in unique_ids)
                fresh=[participant_json(row) for row in con.execute(f"SELECT * FROM participants WHERE id IN ({placeholders})",unique_ids).fetchall()]
            self.json(200,{"results":results,"participants":fresh,"serverTime":utc_now()})
        except Exception:
            con.rollback(); raise
        finally:
            con.close()

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Content-Length","0")
        self.end_headers()


def main():
    init_db()
    server=ThreadingHTTPServer((HOST,PORT),Handler)
    server.daemon_threads=True
    print(f"Imagine Events Desk démarré sur http://{HOST}:{PORT}")
    if USE_POSTGRES:
        print(f"Base PostgreSQL : {redact_database_url()}")
    else:
        print(f"Base locale : {DB_PATH}")
    if db_is_ephemeral():
        print("ATTENTION : cette base est dans un dossier temporaire ; les congrès et les inscrits")
        print("seront perdus au redémarrage. Définissez IMAGINE_DB_PATH vers un dossier persistant.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("Arrêt du serveur…")
    finally:
        server.server_close()

if __name__ == "__main__":
    main()
